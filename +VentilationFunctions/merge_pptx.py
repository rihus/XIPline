"""RH: new file. Cross-platform pptx slide merge for createCombinedVDPReport_RH.m.
Windows merges via PowerPoint COM (Slides.InsertFromFile); macOS/Linux have no
COM and PowerPoint-for-Mac's AppleScript dictionary has no slide-duplicate
command, so slides are recreated here with python-pptx instead (pictures, text
boxes, and tables - matches what Global.exportToPPTX actually produces).

Usage: python merge_pptx.py <output.pptx> <input1.pptx> [input2.pptx ...]
"""
import sys
import io
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.dml import MSO_FILL_TYPE


def copy_text_frame(dest_tf, src_tf):
    dest_tf.word_wrap = src_tf.word_wrap
    for i, para in enumerate(src_tf.paragraphs):
        dest_para = dest_tf.paragraphs[0] if i == 0 else dest_tf.add_paragraph()
        dest_para.alignment = para.alignment
        for run in para.runs:
            r = dest_para.add_run()
            r.text = run.text
            if run.font.size is not None:
                r.font.size = run.font.size
            r.font.bold = run.font.bold
            try:
                if run.font.color and run.font.color.type is not None:
                    r.font.color.rgb = run.font.color.rgb
            except Exception:
                pass


def copy_fill(dest_fill, src_fill):
    # RH: only copy an actual solid fill. src_fill.type is not None for
    # noFill/background too (MSO_FILL_TYPE.BACKGROUND) - calling .solid() for
    # those and then failing to read .fore_color.rgb left destination shapes
    # with an uncoloured solid fill, which renders black instead of the
    # intended "no fill" (transparent/white) background.
    if src_fill.type == MSO_FILL_TYPE.SOLID:
        dest_fill.solid()
        dest_fill.fore_color.rgb = src_fill.fore_color.rgb


def copy_picture(dest_slide, shape):
    image = shape.image
    dest_slide.shapes.add_picture(
        io.BytesIO(image.blob), shape.left, shape.top, shape.width, shape.height)


def copy_textbox(dest_slide, shape):
    tb = dest_slide.shapes.add_textbox(shape.left, shape.top, shape.width, shape.height)
    try:
        copy_fill(tb.fill, shape.fill)
    except Exception:
        pass
    copy_text_frame(tb.text_frame, shape.text_frame)


def copy_table(dest_slide, shape):
    src_table = shape.table
    n_rows = len(src_table.rows)
    n_cols = len(src_table.columns)
    graphic_frame = dest_slide.shapes.add_table(
        n_rows, n_cols, shape.left, shape.top, shape.width, shape.height)
    dest_table = graphic_frame.table

    for r in range(n_rows):
        dest_table.rows[r].height = src_table.rows[r].height
    for c in range(n_cols):
        dest_table.columns[c].width = src_table.columns[c].width

    for r in range(n_rows):
        for c in range(n_cols):
            src_cell = src_table.cell(r, c)
            dest_cell = dest_table.cell(r, c)
            try:
                copy_fill(dest_cell.fill, src_cell.fill)
            except Exception:
                pass
            copy_text_frame(dest_cell.text_frame, src_cell.text_frame)


def merge(output_path, input_paths):
    if not input_paths:
        raise SystemExit("No input files provided")

    src0 = Presentation(input_paths[0])
    prs = Presentation()
    prs.slide_width = src0.slide_width
    prs.slide_height = src0.slide_height
    blank_layout = prs.slide_layouts[6]

    for path in input_paths:
        src = Presentation(path)
        for src_slide in src.slides:
            dest_slide = prs.slides.add_slide(blank_layout)
            for ph in list(dest_slide.placeholders):
                ph._element.getparent().remove(ph._element)
            for shape in src_slide.shapes:
                try:
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        copy_picture(dest_slide, shape)
                    elif shape.has_table:
                        copy_table(dest_slide, shape)
                    elif shape.has_text_frame:
                        copy_textbox(dest_slide, shape)
                    else:
                        print(f"Skipping unsupported shape type on a slide from {path}", file=sys.stderr)
                except Exception as exc:
                    print(f"Skipping shape on a slide from {path}: {exc}", file=sys.stderr)

    prs.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    merge(sys.argv[1], sys.argv[2:])
    print(f"Merged {len(sys.argv) - 2} file(s) into {sys.argv[1]}")
